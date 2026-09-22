#!/usr/bin/env python3
"""Parse the harvested Quálitas filings (tools/qualitas/raw) into the model's data files:

  site/qualitas/data/financials.js   window.Q_FIN  — income statement, balance sheet, cash-flow statement,
                                                     reported ratios and per-quarter KPIs; quarters, YTD and FY
  site/qualitas/data/operations.js   window.Q_OPS  — written premiums by line of business and subsidiary,
                                                     insured units by country/type, solvency, portfolio facts

Sources, in order of precedence for every figure:
  1. The quarterly results report (Informe de resultados, PDF) — CNSF-format income statement for the quarter
     (and the YTD / fiscal-year column), balance sheet, line-of-business and insured-unit tables, solvency,
     portfolio and reported ratios. Each quarter's own report is primary; the comparative column of the
     following year's report fills quarters that have no report of their own.
  2. The SIFIC filing (PDF filed with the BMV/CNBV) — the only source of the cash-flow statement (cumulative
     year-to-date), and a coded cross-check of the balance sheet and income statement.
  3. The company's "Datos Financieros Históricos" workbook — core income-statement and balance-sheet lines,
     ratios, insured units and share count back to 2013, used where no report text is available.
Amounts are stored in THOUSANDS of pesos (the page shows millions); ratios in %.

Usage: python scripts/qualitas/build_data.py
"""
import glob, json, os, re, sys, unicodedata
from datetime import datetime, timezone

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW = os.path.join(ROOT, 'tools', 'qualitas', 'raw')
TXT = os.path.join(RAW, 'text')
OUT_DIR = os.path.join(ROOT, 'site', 'qualitas', 'data')
MANIFEST = os.path.join(RAW, 'manifest.json')
XLSX = os.path.join(RAW, 'DatosFinancierosHistoricos.xlsx')

MONTHS_ES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6, 'julio': 7, 'agosto': 8, 'septiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12}


# ----------------------------------------------------------------------------- helpers
def norm(s):
    s = unicodedata.normalize('NFKD', s or '')
    s = ''.join(c for c in s if not unicodedata.combining(c)).lower()
    s = s.replace('(-)', ' ').replace('*', ' ')
    return re.sub(r'\s+', ' ', s).strip()


# A number is "(1,234)", "-1,234", "$ 1,234", "1,234.5%"; a lone "-" (zero in the CNSF layouts) stays a separate
# token even when the next column's figure follows it after spaces ("Daños -   3,540,000" = [0, 3540000]).
NUM_TOKEN = re.compile(r'\(?(?:-\$\s?|-|\$\s?)?[\d][\d,]*(?:\.\d+)?%?\)?|(?<![\w])-(?![\w])')
NEG_MARK = re.compile(r'\(\s*-\s*\)')  # the "(-)" prefix the CNSF layout prints on deduction lines — not a number


def tokens(line):
    """numeric tokens on a line: returns list of (value, is_pct, raw) — a lone '-' is 0."""
    out = []
    line = NEG_MARK.sub(' ', line)
    for m in NUM_TOKEN.finditer(line):
        raw = m.group(0)
        if raw.strip() == '-':
            out.append((0, False, raw)); continue
        pct = raw.endswith('%') or raw.endswith('%)')
        neg = raw.startswith('(') or raw.startswith('-') or raw.startswith('(-')
        s = re.sub(r'[^\d.]', '', raw)
        if not s or s == '.':
            continue
        v = float(s) if '.' in s else int(s)
        out.append((-v if neg else v, pct, raw))
    return out


def ints(line):
    return [t[0] for t in tokens(line) if not t[1]]


def pages_of(txt):
    parts = re.split(r'=== PAGE (\d+) ===\n', txt)
    out = {}
    for i in range(1, len(parts) - 1, 2):
        out[int(parts[i])] = parts[i + 1]
    return out


def qid_parts(qid):
    fy, q = qid.split('Q'); return int(fy), int(q)


def prev_year_q(qid):
    fy, q = qid_parts(qid); return '%dQ%d' % (fy - 1, q)


def label_of(line):
    """label = text before the first numeric token"""
    line = NEG_MARK.sub(' ', line)
    m = NUM_TOKEN.search(line)
    return norm(line[:m.start()] if m else line)


# ----------------------------------------------------------------------------- statement templates
IS_TPL = [
    ('emitidas', 'written'), ('cedidas', 'ceded'), ('de retencion', 'retained'), ('incremento neto de la reserva', 'reserveInc'), ('riesgos en curso', 'reserveInc'),
    ('primas de retencion devengadas', 'earned'), ('costo neto de adquisicion', 'acqCost'), ('comisiones a agentes', 'commAgents'),
    ('compensaciones adicionales', 'addlComp'), ('reaseguro y reafianzamiento tomado', 'reinsCommTaken'), ('comisiones por reaseguro cedido', 'reinsCommCeded'),
    ('cobertura de exceso', 'xlCover'), ('otros', 'acqOther'), ('otras obligaciones pendientes de cumplir', 'lossCost'),
    ('siniestralidad y otras obligaciones', 'lossGross'), ('siniestralidad recuperada', 'lossRecovered'), ('reclamaciones', 'claims'),
    ('utilidad (perdida) tecnica', 'techResult'), ('incremento neto de otras reservas', 'otherReservesInc'), ('riesgos catastroficos', 'catReserve'),
    ('seguros especializados', 'specReserve'), ('reserva de contingencia', 'contReserve'), ('otras reservas', 'otherReserves'),
    ('operaciones analogas', 'analogous'), ('utilidad (perdida) bruta', 'grossProfit'), ('gastos de operacion netos', 'opex'),
    ('gastos administrativos', 'opexAdmin'), ('remuneraciones', 'opexStaff'), ('depreciaciones', 'opexDA'), ('utilidad (perdida) de la operacion', 'opResult'),
    ('resultado integral de financiamiento', 'rif'), ('de inversiones', 'rifInvestments'), ('por venta de inversiones', 'rifSale'),
    ('por valuacion de inversiones', 'rifValuation'), ('recargo sobre primas', 'rifSurcharge'), ('emision de instrumentos', 'rifDebt'),
    ('reaseguro financiero', 'rifFinReins'), ('intereses por creditos', 'rifLoanInterest'), ('castigos preventivos por importes', 'rifWriteoffReins'),
    ('castigos preventivos por riesgos', 'rifWriteoffCredit'), ('otros', 'rifOther'), ('resultado cambiario', 'rifFx'), ('posicion monetaria', 'rifMonetary'),
    ('resultado de inversiones permanentes', 'associates'), ('antes de impuestos', 'ibt'), ('impuesto', 'tax'), ('antes de operaciones discontinuadas', 'netBeforeDisc'),
    ('operaciones discontinuadas', 'discontinued'), ('utilidad (perdida) del ejercicio', 'netIncome'), ('participacion no controladora', 'nci'), ('participacion controladora', 'netControlling'),
]
BS_A_TPL = [
    ('inversiones', 'inv'), ('valores y operaciones', 'securitiesAndDeriv'), ('valores', 'securities'), ('gubernamentales', 'gov'), ('tasa conocida', 'corpFixed'),
    ('renta variable', 'equities'), ('extranjeros', 'foreign'), ('dividendos por cobrar', 'divReceivable'), ('deterioro de valores', 'impairment'),
    ('dados en prestamo', 'securitiesLent'), ('valores restringidos', 'restricted'), ('operaciones con productos derivados', 'derivatives'), ('deudor por reporto', 'repo'),
    ('cartera de credito (neto)', 'loans'), ('cartera de credito vigente', 'loansPerforming'), ('cartera de credito vencida', 'loansNonPerf'), ('estimaciones preventivas', 'loanProvision'),
    ('inmuebles', 'realEstate'), ('obligaciones laborales', 'laborInv'), ('disponibilidad', 'cash'), ('caja y bancos', 'cashBanks'), ('deudores', 'receivables'),
    ('por primas', 'premRec'), ('subsidio', 'premRecSubsidy'), ('publica federal', 'govRec'), ('primas por cobrar de fianzas', 'bondPremRec'), ('agentes y ajustadores', 'agentsRec'),
    ('documentos por cobrar', 'notesRec'), ('clamaciones', 'bondClaimsRec'), ('otros', 'otherRec'), ('estimacion para castigos', 'recProvision'),
    ('reaseguradores y reafianzadores', 'reinsurers'), ('instituciones de seguros y fianzas', 'reinsInst'), ('depositos retenidos', 'reinsDeposits'),
    ('importes recuperables', 'reinsRecoverable'), ('reaseguradores extranjeros', 'reinsProvision'), ('intermediarios de reaseguro', 'reinsInterm'), ('estimacion para castigos', 'reinsWriteoff'),
    ('inversiones permanentes', 'permInv'), ('subsidiarias', 'permSubs'), ('asociadas', 'permAssoc'), ('otras inversiones permanentes', 'permOther'),
    ('otros activos', 'otherAssets'), ('mobiliario y equipo', 'equipment'), ('activos adjudicados', 'foreclosed'), ('diversos', 'otherAssetsMisc'),
    ('intangibles amortizables', 'intangAmort'), ('intangibles de larga duracion', 'intangLong'), ('suma del activo', 'totalAssets'),
]
BS_L_TPL = [
    ('reservas tecnicas', 'reserves'), ('de riesgos en curso', 'upr'), ('seguros de vida', 'uprLife'), ('accidentes y enfermedades', 'uprAccident'), ('seguros de danos', 'uprPC'),
    ('reafianzamiento tomado', 'uprReinsTaken'), ('de fianzas en vigor', 'bondsInForce'), ('obligaciones pendientes de cumplir', 'claimsReserves'), ('polizas vencidas', 'claimsPending'),
    ('no reportados', 'ibnr'), ('fondos en administracion', 'fundsAdmin'), ('primas en deposito', 'premDeposit'), ('reserva de contingencia', 'contReserveL'),
    ('seguros especializados', 'specReserveL'), ('riesgos catastroficos', 'catReserveL'), ('obligaciones laborales', 'laborReserves'), ('acreedores', 'creditors'),
    ('agentes y ajustadores', 'creditorsAgents'), ('administracion de perdidas', 'lossFunds'), ('responsabilidades de fianzas', 'bondCreditors'), ('diversos', 'creditorsMisc'),
    ('reaseguradores y reafianzadores', 'reinsPayable'), ('instituciones de seguros y fianzas', 'reinsPayInst'), ('depositos retenidos', 'reinsPayDeposits'),
    ('otras participaciones', 'reinsPayOther'), ('intermediarios', 'reinsPayInterm'), ('productos derivados', 'derivLiab'), ('financiamientos obtenidos', 'financing'),
    ('emision de deuda', 'debtIssued'), ('no susceptibles', 'subDebt'), ('otros titulos', 'otherDebt'), ('reaseguro financiero', 'finReins'), ('otros pasivos', 'otherLiab'),
    ('trabajadores en la utilidad', 'ptuProv'), ('pago de impuestos', 'taxProv'), ('otras obligaciones', 'otherObl'), ('creditos diferidos', 'deferredCredits'),
    ('suma del pasivo', 'totalLiab'), ('fondo social pagado', 'paidCapital'), ('capital o fondo social', 'socialCapital'), ('no suscrito', 'unsubscribed'),
    ('no exhibido', 'unpaid'), ('acciones propias recompradas', 'treasury'), ('conversion obligatoria', 'convSub'), ('capital ganado', 'earnedCapital'),
    ('reservas', 'capReserves'), ('legal', 'legalReserve'), ('adquisicion de acciones propias', 'buybackReserve'), ('otras', 'otherCapReserves'),
    ('superavit por valuacion', 'valuationSurplus'), ('inversiones permanentes', 'permInvEq'), ('ejercicios anteriores', 'retained'), ('del ejercicio', 'netIncomeEq'),
    ('activos no monetarios', 'nonMonetary'), ('participacion controladora', 'controllingEquity'), ('participacion no controladora', 'nci'),
    ('suma del capital', 'totalEquity'), ('suma del pasivo y capital', 'totalLiabEquity'),
]
# "Reserva para obligaciones pendientes de cumplir" and its two sub-lines print without a label in the report PDFs
BS_L_UNLABELED = ['claimsReserves', 'claimsPending', 'ibnr']

SIFIC_IS = {'410': 'written', '420': 'ceded', '430': 'retained', '440': 'reserveInc', '450': 'earned', '460': 'acqCost', '470': 'commAgents', '480': 'addlComp', '490': 'reinsCommTaken',
            '500': 'reinsCommCeded', '510': 'xlCover', '520': 'acqOther', '530': 'lossCost', '540': 'lossGross', '550': 'lossRecovered', '560': 'claims', '570': 'techResult',
            '580': 'otherReservesInc', '625': 'analogous', '630': 'grossProfit', '640': 'opex', '650': 'opexAdmin', '660': 'opexStaff', '670': 'opexDA', '680': 'opResult',
            '690': 'rif', '700': 'rifInvestments', '710': 'rifSale', '720': 'rifValuation', '730': 'rifSurcharge', '750': 'rifDebt', '760': 'rifFinReins', '820': 'rifLoanInterest',
            '830': 'rifWriteoffReins', '840': 'rifWriteoffCredit', '770': 'rifOther', '780': 'rifFx', '790': 'rifMonetary', '795': 'associates', '801': 'ibt', '802': 'tax',
            '804': 'netBeforeDisc', '808': 'discontinued', '805': 'netIncome', '806': 'nci', '807': 'netControlling'}
SIFIC_BS = {'100': 'totalAssets', '110': 'inv', '111': 'securitiesAndDeriv', '112': 'securities', '113': 'gov', '115': 'corpFixed', '116': 'equities', '117': 'foreign', '120': 'divReceivable',
            '121': 'impairment', '123': 'securitiesLent', '124': 'restricted', '125': 'derivatives', '126': 'repo', '171': 'loans', '172': 'loansPerforming', '133': 'loansNonPerf',
            '135': 'loanProvision', '137': 'realEstate', '140': 'laborInv', '141': 'cash', '142': 'cashBanks', '143': 'receivables', '144': 'premRec', '173': 'premRecSubsidy', '174': 'govRec',
            '145': 'agentsRec', '146': 'notesRec', '175': 'bondClaimsRec', '148': 'otherRec', '149': 'recProvision', '150': 'reinsurers', '151': 'reinsInst', '152': 'reinsDeposits',
            '176': 'reinsRecoverable', '177': 'reinsProvision', '156': 'reinsInterm', '158': 'reinsWriteoff', '159': 'permInv', '160': 'permSubs', '161': 'permAssoc', '162': 'permOther',
            '163': 'otherAssets', '164': 'equipment', '165': 'foreclosed', '166': 'otherAssetsMisc', '178': 'intangAmort', '179': 'intangLong', '200': 'totalLiab', '210': 'reserves',
            '211': 'upr', '212': 'uprLife', '213': 'uprAccident', '214': 'uprPC', '249': 'uprReinsTaken', '215': 'bondsInForce', '216': 'claimsReserves', '217': 'claimsPending', '218': 'ibnr',
            '220': 'fundsAdmin', '221': 'premDeposit', '225': 'contReserveL', '226': 'specReserveL', '224': 'catReserveL', '227': 'laborReserves', '228': 'creditors', '229': 'creditorsAgents',
            '230': 'lossFunds', '231': 'bondCreditors', '232': 'creditorsMisc', '233': 'reinsPayable', '234': 'reinsPayInst', '235': 'reinsPayDeposits', '236': 'reinsPayOther', '237': 'reinsPayInterm',
            '238': 'derivLiab', '239': 'financing', '240': 'debtIssued', '241': 'subDebt', '242': 'otherDebt', '243': 'finReins', '244': 'otherLiab', '245': 'ptuProv', '246': 'taxProv',
            '247': 'otherObl', '248': 'deferredCredits', '300': 'totalEquity', '301': 'contributedCapital', '310': 'paidCapital', '311': 'socialCapital', '312': 'unsubscribed', '313': 'unpaid',
            '314': 'treasury', '315': 'convSub', '302': 'earnedCapital', '316': 'capReserves', '317': 'legalReserve', '318': 'buybackReserve', '319': 'otherCapReserves', '320': 'valuationSurplus',
            '321': 'permInvEq', '323': 'retained', '324': 'netIncomeEq', '325': 'nonMonetary', '303': 'controllingEquity', '326': 'nci', '327': 'totalLiabEquity'}
SIFIC_CF = {'21100': 'netIncome', '21200': 'nonCashAdj', '21201': 'valuation', '21202': 'badDebt', '21203': 'impairment', '21204': 'da', '21205': 'reservesAdj', '21206': 'provisions',
            '21207': 'taxes', '21208': 'associates', '21209': 'discontinued', '21301': 'marginAccounts', '21302': 'chgSecurities', '21303': 'chgRepo', '21304': 'chgSecLending',
            '21305': 'chgDerivAssets', '21306': 'chgPremRec', '21307': 'chgReceivables', '21308': 'chgReinsurers', '21309': 'chgForeclosed', '21310': 'chgOtherOpAssets',
            '21311': 'chgClaimsObligations', '21314': 'chgDerivLiab', '21315': 'chgOtherOpLiab', '21316': 'chgHedges', '21300': 'cfo', '21401': 'saleProceedsPPE', '21402': 'capexPPE',
            '21403': 'saleSubs', '21404': 'acqSubs', '21405': 'salePermInv', '21406': 'acqPermInv', '21407': 'divReceived', '21408': 'intangibles', '21409': 'saleLongAssets',
            '21410': 'saleOtherLong', '21411': 'acqOtherLong', '21400': 'cfi', '21501': 'sharesIssued', '21502': 'capitalReimb', '21503': 'dividendsPaid', '21504': 'buybacks',
            '21505': 'subDebtIssued', '21506': 'subDebtPaid', '21500': 'cff', '21000': 'netChangeCash', '22000': 'fxEffectCash', '23000': 'cashBegin', '20000': 'cashEnd'}


def parse_template(lines, tpl, unlabeled=None, start_key=None):
    """Sequential label matching: walk the lines; each numeric line is matched to the next template entry
    whose pattern is contained in the line's label (longest pattern within a 7-entry window wins)."""
    out = {}
    p = 0
    pending_unlabeled = list(unlabeled or [])
    for line in lines:
        tk = tokens(line)
        if len(tk) < 2:
            continue
        lab = label_of(line)
        vals = [t[0] for t in tk[:2]]
        if not lab and pending_unlabeled and p > 0:
            k = pending_unlabeled.pop(0)
            out[k] = vals; continue
        best = None
        for j in range(p, min(p + 7, len(tpl))):
            pat, key = tpl[j]
            if pat in lab and (best is None or len(pat) > len(tpl[best][0])):
                best = j
        if best is None:
            continue
        out[tpl[best][1]] = vals
        p = best + 1
    return out


def k_thousands(v):
    return None if v is None else round(v / 1000.0, 3)


# ----------------------------------------------------------------------------- IR quarterly report
def parse_report(qid, txt):
    """returns dict with is_q, is_ytd (or fy), bs, ops..., each value = [current, prior-year]"""
    P = pages_of(txt)
    fy, q = qid_parts(qid)
    res = {'qid': qid, 'fy': fy, 'q': q}
    m = re.search(r'SOURCE:\s*(\S+)', txt)
    res['url'] = m.group(1) if m else None
    for n, t in sorted(P.items()):
        m = re.search(r'Ciudad de M.xico,\s*(\d{1,2})\s*de\s*(\w+)\s*de\s*(\d{4})', t)
        if m and 'date' not in res:
            mo = MONTHS_ES.get(norm(m.group(2)))
            if mo:
                res['date'] = '%s-%02d-%02d' % (m.group(3), mo, int(m.group(1)))
    is_pages = []
    for n, t in sorted(P.items()):
        lines = [l.strip() for l in t.split('\n')]
        if any(l.startswith('Suma del Activo') for l in lines):
            res['bs_a'] = parse_template(lines, BS_A_TPL)
        if any(l.startswith('Suma del Pasivo') for l in lines):
            res['bs_l'] = parse_template(lines, BS_L_TPL, unlabeled=BS_L_UNLABELED)
        if any(re.match(r'^Emitidas\s+[\d(]', l) for l in lines) and any('del Ejercicio' in l for l in lines):
            is_pages.append(parse_template(lines, IS_TPL))
    if is_pages:
        res['is_q'] = is_pages[0]
        if len(is_pages) > 1:
            res['is_ytd'] = is_pages[1]
    # balance sheet merged
    bs = {}
    for part in ('bs_a', 'bs_l'):
        for k, v in res.get(part, {}).items():
            bs[k] = v
    if bs:
        res['bs'] = bs
    # ---- operating tables
    full = '\n'.join(P[n] for n in sorted(P))
    ops = {}
    seg_pat = [('tradicional', 'trad'), ('individual', 'ind'), ('flotillas', 'fleet'), ('instituciones financieras', 'fin'), ('subsidiarias en el extranjero', 'intl'), ('total', 'total')]
    for n, t in sorted(P.items()):
        if 'nea de negocio' in norm(t) and 'flotillas' in norm(t):
            seg = {}
            for l in t.split('\n'):
                lab = label_of(l)
                for pat, key in seg_pat:
                    if lab.startswith(pat) and key not in seg:
                        v = ints(l)
                        if len(v) >= 2:
                            seg[key] = {'q': [v[0], v[1]]}
                            if len(v) >= 4:
                                seg[key]['ytd'] = [v[2], v[3]]
                        break
            if seg:
                ops['segments'] = seg; break
    unit_pat = [('mexico', 'mx'), ('automoviles', 'autos'), ('camiones', 'trucks'), ('motocicletas', 'motos'), ('el salvador', 'sv'), ('costa rica', 'cr'), ('estados unidos', 'us'), ('peru', 'pe'), ('colombia', 'co'), ('unidades aseguradas', 'total'), ('total', 'total')]
    for n, t in sorted(P.items()):
        nt = norm(t)
        if 'unidades aseguradas' in nt and 'camiones' in nt and 'costa rica' in nt:
            un = {}
            for l in t.split('\n'):
                lab = label_of(l)
                if not lab or re.search(r'\d\s*T\s*\d\d', l) or re.search(r'\dQ\d\d', l):
                    continue  # header rows such as "Unidades aseguradas 2T26 1T26 Δ% 2T25 Δ%"
                for pat, key in unit_pat:
                    if lab.startswith(pat) and key not in un:
                        v = [x for x in ints(l) if isinstance(x, int) or float(x).is_integer()]
                        v = [int(x) for x in v]
                        if len(v) >= 2 and v[0] < 100000 and (key != 'total' or v[0] > 1000):
                            un[key] = {'cur': v[0], 'prevQ': v[1] if len(v) >= 3 else None, 'prevY': v[2] if len(v) >= 3 else v[1]}
                        break
            if un:
                ops['units'] = un; break
    sub_pat = [('q es', 'es'), ('q cr', 'cr'), ('q ic', 'ic'), ('q p', 'pe'), ('q col', 'co'), ('verticales', 'verticals'), ('total', 'total')]
    for n, t in sorted(P.items()):
        nt = norm(t)
        if 'q cr' in nt and 'q es' in nt and ('verticales' in nt or 'q ic' in nt):
            sb = {}
            for l in t.split('\n'):
                lab = label_of(l)
                for pat, key in sub_pat:
                    if lab.startswith(pat) and key not in sb and not (pat == 'q p' and lab.startswith('q pe')):
                        v = ints(l)
                        if len(v) >= 2:
                            sb[key] = {'q': [v[0], v[1]]}
                            if len(v) >= 4:
                                sb[key]['ytd'] = [v[2], v[3]]
                        break
            if 'es' in sb or 'cr' in sb:
                ops['subsidiaries'] = sb; break
    m = re.search(r'requerimiento de capital regulatorio se situ. en \$?\s*([\d,]+)\s*millones.{0,120}?margen de\s*solvencia de \$?\s*([\d,]+)\s*millones', norm(full).replace('\n', ' '), re.S)
    if m:
        ops['solvency'] = {'rcs': int(m.group(1).replace(',', '')), 'margin': int(m.group(2).replace(',', ''))}
        m2 = re.search(r'ndice de (?:margen de )?solvencia\s*(?:de)?\s*(\d{3})\s*%', norm(full).replace('\n', ' '))
        if m2:
            ops['solvency']['index'] = int(m2.group(1))
        else:
            ops['solvency']['index'] = round(100 * (ops['solvency']['rcs'] + ops['solvency']['margin']) / ops['solvency']['rcs'])
    nf = norm(full).replace('\n', ' ')
    m = re.search(r'(\d{2}(?:\.\d+)?)\s*%\s*(?:del portafolio se encuentra )?invertido en renta fija', nf)
    if m:
        ops.setdefault('portfolio', {})['fiPct'] = float(m.group(1))
    m = re.search(r'duraci.n (?:total )?(?:de nuestro portafolio de inversi.n se posicion. en |de )?(\d(?:\.\d+)?)\s*a.os', nf) or re.search(r'duraci.n[^.]{0,80}?(\d\.\d+)\s*a.os', nf)
    if m:
        ops.setdefault('portfolio', {})['duration'] = float(m.group(1))
    # "Cifras relevantes" page: reported ratios (first number = quarter, then prior-year quarter; YTD pair after)
    for n, t in sorted(P.items()):
        if 'cifras relevantes' in norm(t) or ('estado de resultados' in norm(t) and 'roe 12m' in norm(t)):
            rep = {}
            for l in t.split('\n'):
                lab = label_of(l)
                tk = [x[0] for x in tokens(l) if x[1]]  # percentages
                allv = [x[0] for x in tokens(l)]
                if lab.startswith('indice combinado ajustado') and tk:
                    rep['combinedAdj'] = tk[0]; rep['combinedAdjPrev'] = tk[1] if len(tk) > 1 else None
                    if len(tk) >= 4: rep['combinedAdjYtd'] = tk[2]
                elif lab.startswith('rendimiento sobre inversiones') and tk:
                    rep['rsi'] = tk[0]; rep['rsiPrev'] = tk[1] if len(tk) > 1 else None
                    if len(tk) >= 4: rep['rsiYtd'] = tk[2]
                elif lab.startswith('roe del periodo') and tk:
                    rep['roePeriod'] = tk[0]; rep['roePeriodPrev'] = tk[1] if len(tk) > 1 else None
                elif lab.startswith('roe 12m') and tk:
                    rep['roe12'] = tk[0]; rep['roe12Prev'] = tk[1] if len(tk) > 1 else None
                elif lab.startswith('activos invertidos') and allv:
                    rep['float'] = allv[0]; rep['floatPrev'] = allv[1] if len(allv) > 1 else None
                elif lab.startswith('inversiones total') and allv:
                    rep['invTotal'] = allv[0]; rep['invTotalPrev'] = allv[1] if len(allv) > 1 else None
            if rep:
                ops['reported'] = rep; break
    res['ops'] = ops
    return res


# ----------------------------------------------------------------------------- SIFIC filing
def parse_sific(qid, txt):
    P = pages_of(txt)
    out = {'qid': qid, 'bs': {}, 'is_ytd': {}, 'cf_ytd': {}}
    m = re.search(r'SOURCE:\s*(\S+)', txt)
    out['url'] = m.group(1) if m else None
    NUM = r'-?\(?[\d,]+\)?'
    for n, t in sorted(P.items()):
        head = t[:400].upper()
        if 'FLUJOS DE EFECTIVO' in head: kind, mp = 'cf_ytd', SIFIC_CF
        elif 'ESTADO DE RESULTADOS' in head: kind, mp = 'is_ytd', SIFIC_IS
        elif 'BALANCE GENERAL' in head: kind, mp = 'bs', SIFIC_BS
        else: continue
        prev_code = None
        for line in t.split('\n'):
            line = line.strip()
            m1 = re.match(r'^(\d{3,5})\s+(' + NUM + r')\s+(' + NUM + r')\s*(\D.*)?$', line)
            m2 = re.match(r'^(\d{3,5})\s+(.*?)\s+(' + NUM + r')\s+(' + NUM + r')\s*$', line)
            m3 = re.match(r'^(' + NUM + r')\s+(' + NUM + r')\s*$', line)  # numbers-only continuation of a wrapped label
            if m1:
                code, a, b = m1.group(1), m1.group(2), m1.group(3)
            elif m2:
                code, a, b = m2.group(1), m2.group(3), m2.group(4)
            elif m3 and prev_code and prev_code not in out[kind]:
                code, a, b = prev_code, m3.group(1), m3.group(2)
            else:
                mm = re.match(r'^(\d{3,5})\s+\D', line)
                if mm: prev_code = mm.group(1)
                continue
            key = mp.get(code)
            if key and key not in out[kind]:
                out[kind][key] = [tokens(a)[0][0], tokens(b)[0][0]]
            prev_code = None
    return out


# ----------------------------------------------------------------------------- historical workbook
XLSX_ROWS = {'prima emitida': 'written', 'prima retenida': 'retained', 'prima devengada': 'earned', 'costo de adquisicion': 'acqCost', 'costo de siniestralidad': 'lossCost',
             'resultado tecnico': 'techResult', 'gastos de operacion': 'opex', 'resultado operativo': 'opResult', 'resultado integral de financiamiento': 'rif', 'impuestos': 'tax',
             'resultado neto': 'netIncome', 'inversiones': 'inv', 'activo total': 'totalAssets', 'reservas tecnicas': 'reserves', 'pasivo total': 'totalLiab', 'capital contable': 'totalEquity',
             'indice de adquisicion': 'acqRatioRep', 'indice de siniestralidad': 'lossRatioRep', 'indice de operacion': 'opRatioRep', 'indice combinado': 'combinedRep',
             'rendimiento sobre las inversiones': 'rsi', 'roe 12m': 'roe12', 'unidades aseguradas': 'units', 'total de acciones en circulacion': 'shares'}


def parse_xlsx(path):
    if not os.path.exists(path):
        return {}
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Financieros'] if 'Financieros' in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = None
    for r in rows:
        if r and r[0] and norm(str(r[0])).startswith('estado de resultados'):
            header = r; break
    if not header:
        return {}
    cols = {}
    for j, h in enumerate(header):
        if j == 0 or h is None: continue
        h = str(h).strip()
        m = re.match(r'^([1-4])T(\d{2})$', h)
        if m: cols[j] = '20%sQ%s' % (m.group(2), m.group(1))
        elif re.match(r'^\d{4}$', h): cols[j] = 'FY' + h
    data = {}
    for r in rows:
        if not r or not r[0]: continue
        lab = norm(str(r[0]))
        key = None
        for pat, k in XLSX_ROWS.items():
            if lab.startswith(pat): key = k; break
        if not key: continue
        for j, pid in cols.items():
            v = r[j] if j < len(r) else None
            if v is None or v == '': continue
            try: v = float(v)
            except (TypeError, ValueError): continue
            data.setdefault(pid, {})[key] = v
    return data


# ----------------------------------------------------------------------------- assembly
BUILD_LOG = os.path.join(RAW, 'build-log.json')
LOG = []  # structural parse warnings, written to tools/qualitas/raw/build-log.json and shown on quality.html


def logw(file, msg):
    LOG.append({'file': file, 'msg': msg}); print('WARN %s: %s' % (file, msg))


def main():
    man = json.load(open(MANIFEST, encoding='utf-8')) if os.path.exists(MANIFEST) else {'items': {}}
    reports, sifics = {}, {}
    for f in sorted(glob.glob(os.path.join(TXT, 'reports', '*.txt'))):
        qid = os.path.basename(f)[:-4]
        r = reports[qid] = parse_report(qid, open(f, encoding='utf-8').read())
        name = 'reports/%s.txt' % qid
        if not r.get('is_q'): logw(name, 'no quarterly income statement found')
        elif any(r['is_q'].get(k) is None for k in ('written', 'earned', 'lossCost', 'netIncome')): logw(name, 'income statement missing a core line (written / earned / claims / net income)')
        if not r.get('bs'): logw(name, 'no balance sheet found')
        elif any(r['bs'].get(k) is None for k in ('totalAssets', 'totalLiab', 'totalEquity')): logw(name, 'balance sheet missing a total')
        if not r.get('date'): logw(name, 'release date not found')
        o = r.get('ops', {})
        if not o.get('units'): logw(name, 'insured-units table not found')
        if not o.get('segments'): logw(name, 'line-of-business table not found')
        if not o.get('solvency'): logw(name, 'solvency paragraph not found')
    for f in sorted(glob.glob(os.path.join(TXT, 'sific', '*.txt'))):
        qid = os.path.basename(f)[:-4]
        txt = open(f, encoding='utf-8').read()
        name = 'sific/%s.txt' % qid
        if len(txt) < 2000:
            logw(name, 'image-only PDF (no text layer), skipped'); continue
        s = sifics[qid] = parse_sific(qid, txt)
        if not s['cf_ytd'] or s['cf_ytd'].get('cashEnd') is None: logw(name, 'cash-flow statement not parsed')
        if not s['bs'] or s['bs'].get('totalAssets') is None: logw(name, 'balance sheet not parsed')
        if not s['is_ytd'] or s['is_ytd'].get('written') is None: logw(name, 'income statement not parsed')
    xl = parse_xlsx(XLSX)
    if not xl: logw('DatosFinancierosHistoricos.xlsx', 'historical workbook missing or unreadable (pre-2019 quarters will be empty)')

    Q, YTD, FY = {}, {}, {}
    def qrec(qid):
        fy, q = qid_parts(qid)
        return Q.setdefault(qid, {'id': qid, 'fy': fy, 'q': q, 'label': '%dQ%s' % (q, str(fy)[2:]), 'is': {}, 'bs': {}, 'cf': {}, 'kpi': {}, 'sources': {}, 'shares': {}})
    def yrec(fy, months):
        return YTD.setdefault('%dM%d' % (fy, months), {'id': '%dM%d' % (fy, months), 'fy': fy, 'months': months, 'is': {}, 'cf': {}, 'kpi': {}, 'sources': {}})
    def frec(fy):
        return FY.setdefault('FY%d' % fy, {'id': 'FY%d' % fy, 'fy': fy, 'is': {}, 'bs': {}, 'cf': {}, 'kpi': {}, 'sources': {}})

    # Source precedence (lower rank wins): 1 = the period's own IR report, 2 = the period's own SIFIC filing
    # (or a difference of two own filings), 3 = the prior-year comparative column of the following year's
    # report, 4 = the comparative column of a SIFIC filing, 9 = the IR historical workbook (rounded millions).
    # Precedence is applied per statement, never line by line, so figures from two filings are not mixed.
    R_REPORT, R_SIFIC, R_REPORT_CMP, R_SIFIC_CMP, R_XLSX = 1, 2, 3, 4, 9

    def put(target, key_vals, col, rank, src):
        """key_vals: {key:[cur, prev]}; col 0 = current period, 1 = prior-year comparative"""
        vals = {k: v[col] for k, v in key_vals.items() if len(v) > col and v[col] is not None}
        if not vals: return
        cur = target.get('_rank')
        if cur is not None and rank > cur: return
        if cur is None or rank < cur:
            for k in [k for k in target if not k.startswith('_')]: del target[k]
            target['_src'] = src; target['_rank'] = rank
        for k, val in vals.items():
            if k not in target:
                target[k] = k_thousands(val) if isinstance(val, (int, float)) else val

    # 1) quarterly reports — own quarter is primary, comparative fills the prior year
    for qid in sorted(reports):
        r = reports[qid]; fy, q = r['fy'], r['q']
        src = {'url': r.get('url'), 'date': r.get('date'), 'title': 'Informe de resultados %dT%s' % (q, str(fy)[2:]), 'primary': True}
        srcc = dict(src, primary=False, comparative=True)
        if r.get('is_q'):
            put(qrec(qid)['is'], r['is_q'], 0, R_REPORT, src)
            put(qrec(prev_year_q(qid))['is'], r['is_q'], 1, R_REPORT_CMP, srcc)
        if r.get('bs'):
            put(qrec(qid)['bs'], r['bs'], 0, R_REPORT, src)
            put(qrec(prev_year_q(qid))['bs'], r['bs'], 1, R_REPORT_CMP, srcc)
        if r.get('is_ytd'):
            if q == 4:
                put(frec(fy)['is'], r['is_ytd'], 0, R_REPORT, src); put(frec(fy - 1)['is'], r['is_ytd'], 1, R_REPORT_CMP, srcc)
                put(yrec(fy, 12)['is'], r['is_ytd'], 0, R_REPORT, src); put(yrec(fy - 1, 12)['is'], r['is_ytd'], 1, R_REPORT_CMP, srcc)
            else:
                put(yrec(fy, 3 * q)['is'], r['is_ytd'], 0, R_REPORT, src); put(yrec(fy - 1, 3 * q)['is'], r['is_ytd'], 1, R_REPORT_CMP, srcc)
        if q == 1 and r.get('is_q'):
            put(yrec(fy, 3)['is'], r['is_q'], 0, R_REPORT, src); put(yrec(fy - 1, 3)['is'], r['is_q'], 1, R_REPORT_CMP, srcc)
    # 2) SIFIC — cash flow (YTD), balance sheet + income statement for periods without a report
    for qid in sorted(sifics):
        s = sifics[qid]; fy, q = qid_parts(qid)
        src = {'url': s.get('url'), 'title': 'Reporte SIFIC %dT%s (BMV/CNBV)' % (q, str(fy)[2:]), 'primary': True}
        srcc = dict(src, primary=False, comparative=True)
        if s['cf_ytd']:
            put(yrec(fy, 3 * q)['cf'], s['cf_ytd'], 0, R_SIFIC, src); put(yrec(fy - 1, 3 * q)['cf'], s['cf_ytd'], 1, R_SIFIC_CMP, srcc)
            if q == 4:
                put(frec(fy)['cf'], s['cf_ytd'], 0, R_SIFIC, src); put(frec(fy - 1)['cf'], s['cf_ytd'], 1, R_SIFIC_CMP, srcc)
            if q == 1:
                put(qrec(qid)['cf'], s['cf_ytd'], 0, R_SIFIC, src); put(qrec(prev_year_q(qid))['cf'], s['cf_ytd'], 1, R_SIFIC_CMP, srcc)
        if s['bs']:
            put(qrec(qid)['bs'], s['bs'], 0, R_SIFIC, src); put(qrec(prev_year_q(qid))['bs'], s['bs'], 1, R_SIFIC_CMP, srcc)
        if s['is_ytd']:
            if q == 4:
                put(frec(fy)['is'], s['is_ytd'], 0, R_SIFIC, src); put(frec(fy - 1)['is'], s['is_ytd'], 1, R_SIFIC_CMP, srcc)
                put(yrec(fy, 12)['is'], s['is_ytd'], 0, R_SIFIC, src); put(yrec(fy - 1, 12)['is'], s['is_ytd'], 1, R_SIFIC_CMP, srcc)
            else:
                put(yrec(fy, 3 * q)['is'], s['is_ytd'], 0, R_SIFIC, src); put(yrec(fy - 1, 3 * q)['is'], s['is_ytd'], 1, R_SIFIC_CMP, srcc)
            if q == 1:
                put(qrec(qid)['is'], s['is_ytd'], 0, R_SIFIC, src); put(qrec(prev_year_q(qid))['is'], s['is_ytd'], 1, R_SIFIC_CMP, srcc)
    # 2b) the CNSF cash-flow layout: "Flujos netos de efectivo de actividades de operación" (code 21300) is the
    # sum of the changes in operating items only; net income and the non-cash adjustments sit above it. Keep
    # that subtotal as `opChanges` and present `cfo` the conventional way (net income + adjustments + changes)
    # so that CFO + CFI + CFF = net change in cash.
    def cf_total(c):
        if c.get('_cfoTotal') or c.get('cfo') is None or c.get('netIncome') is None or c.get('nonCashAdj') is None: return
        c['opChanges'] = c['cfo']; c['cfo'] = round(c['netIncome'] + c['nonCashAdj'] + c['opChanges'], 3); c['_cfoTotal'] = True
    for r in list(YTD.values()) + list(FY.values()) + list(Q.values()): cf_total(r['cf'])
    # 2c) quarters without an own report or filing: income statement = YTD(q) − YTD(q−1) of two own filings
    def diff_dict(a, b):
        o = {}
        for k, v in a.items():
            if k.startswith('_'): continue
            if isinstance(v, (int, float)) and isinstance(b.get(k), (int, float)): o[k] = round(v - b[k], 3)
        return o
    for y in list(YTD.values()):
        m = y['months']
        if m == 3 or y['is'].get('_rank', 99) > R_SIFIC: continue
        b = YTD.get('%dM%d' % (y['fy'], m - 3))
        if not b or b['is'].get('_rank', 99) > R_SIFIC: continue
        rec = qrec('%dQ%d' % (y['fy'], m // 3))
        if rec['is'].get('_rank', 99) <= R_SIFIC: continue
        rec['is'] = dict(diff_dict(y['is'], b['is']), _rank=R_SIFIC, _src=dict(y['is']['_src'], derived='difference of two year-to-date filings'))
    # 3) workbook: core lines for every quarter/year since 2013 (only for periods with no filing at all)
    for pid, vals in xl.items():
        if pid.startswith('FY'):
            tgt = frec(int(pid[2:]))
        else:
            tgt = qrec(pid)
        xsrc = {'url': (man['items'].get('xlsx:historicos') or {}).get('url'), 'title': 'Datos Financieros Históricos (IR workbook)', 'workbook': True}
        for k, v in vals.items():
            if k in ('acqRatioRep', 'lossRatioRep', 'opRatioRep', 'combinedRep', 'rsi', 'roe12'):
                if v is not None and v < 5:
                    tgt['kpi'].setdefault(k, round(100 * v, 2))
            elif k == 'units':
                tgt['kpi'].setdefault('units', round(v / 1000.0, 3))
            elif k == 'shares':
                tgt.setdefault('shares', {}).setdefault('current', int(v))
            elif k in ('inv', 'totalAssets', 'reserves', 'totalLiab', 'totalEquity'):
                if '_rank' not in tgt['bs'] or tgt['bs']['_rank'] == R_XLSX:
                    tgt['bs'].setdefault('_rank', R_XLSX); tgt['bs'].setdefault('_src', xsrc); tgt['bs'].setdefault(k, round(v * 1000, 3))
            else:
                if '_rank' not in tgt['is'] or tgt['is']['_rank'] == R_XLSX:
                    tgt['is'].setdefault('_rank', R_XLSX); tgt['is'].setdefault('_src', xsrc); tgt['is'].setdefault(k, round(v * 1000, 3))
        tgt['sources'].setdefault('xlsx', xsrc)
    # 4) derived: quarterly cash flow = YTD(q) − YTD(q−1)
    for qid, rec in Q.items():
        fy, q = rec['fy'], rec['q']
        if not rec['cf'] and q > 1:
            a, b = YTD.get('%dM%d' % (fy, 3 * q)), YTD.get('%dM%d' % (fy, 3 * (q - 1)))
            if a and a['cf'] and b and b['cf']:
                for k, v in a['cf'].items():
                    if k.startswith('_'): continue
                    if k in ('cashBegin',):
                        rec['cf'][k] = b['cf'].get('cashEnd')
                    elif k == 'cashEnd':
                        rec['cf'][k] = v
                    elif k in b['cf'] and isinstance(v, (int, float)):
                        rec['cf'][k] = round(v - b['cf'][k], 3)
                rec['cf']['_derived'] = True; rec['cf']['_cfoTotal'] = True
                rec['sources']['cf'] = dict(a['cf'].get('_src', {}), derived='YTD difference of two SIFIC filings')
        elif rec['cf'] and '_src' in rec['cf']:
            rec['sources']['cf'] = rec['cf']['_src']
        for part in ('is', 'bs'):
            if rec[part].get('_src'): rec['sources'][part] = rec[part]['_src']
    # 5) ratios and per-share KPIs
    def ratios(rec, months=3):
        is_ = rec['is']; k = rec['kpi']
        def g(x): return is_.get(x)
        if g('written') and g('retained') and g('earned') is not None:
            if g('acqCost') is not None and g('retained'): k['acqRatio'] = round(100 * g('acqCost') / g('retained'), 2)
            if g('lossCost') is not None and g('earned'): k['lossRatio'] = round(100 * g('lossCost') / g('earned'), 2)
            if g('opex') is not None and g('written'): k['opRatio'] = round(100 * g('opex') / g('written'), 2)
            if all(x in k for x in ('acqRatio', 'lossRatio', 'opRatio')): k['combined'] = round(k['acqRatio'] + k['lossRatio'] + k['opRatio'], 2)
            if g('acqCost') is not None and g('lossCost') is not None and g('opex') is not None and g('earned'): k['combinedAdj'] = round(100 * (g('acqCost') + g('lossCost') + g('opex')) / g('earned'), 2)
            if g('techResult') is not None and g('earned'): k['techMargin'] = round(100 * g('techResult') / g('earned'), 2)
            if g('opResult') is not None and g('earned'): k['opMargin'] = round(100 * g('opResult') / g('earned'), 2)
            if g('netIncome') is not None and g('written'): k['netMargin'] = round(100 * g('netIncome') / g('written'), 2)
            if g('retained') and g('written'): k['retention'] = round(100 * g('retained') / g('written'), 2)
        if g('ibt') and g('tax') is not None: k['taxRate'] = round(100 * g('tax') / g('ibt'), 2)
        sh = (rec.get('shares') or {}).get('current')
        if g('netIncome') is not None and sh: k['eps'] = round(g('netIncome') * 1000 / sh, 4)
    for rec in Q.values(): ratios(rec)
    for rec in YTD.values(): ratios(rec)
    for rec in FY.values(): ratios(rec)
    # reported ratios / solvency / portfolio from the reports into kpi, ops into operations
    OPS = []
    for qid, r in sorted(reports.items()):
        rec = qrec(qid); o = r.get('ops', {})
        rep = o.get('reported', {})
        for k in ('rsi', 'roe12', 'roePeriod', 'combinedAdj', 'float', 'invTotal'):
            if rep.get(k) is not None:
                rec['kpi'][k if k != 'combinedAdj' else 'combinedAdjRep'] = rep[k]
        if rep.get('rsiYtd') is not None:
            y = YTD.get('%dM%d' % (r['fy'], 3 * r['q']))
            if y: y['kpi']['rsi'] = rep['rsiYtd']
        if o.get('solvency'):
            rec['kpi']['rcs'] = o['solvency']['rcs'] * 1000; rec['kpi']['solvMargin'] = o['solvency']['margin'] * 1000; rec['kpi']['solvIndex'] = o['solvency'].get('index')
        if o.get('portfolio'):
            rec['kpi']['fiPct'] = o['portfolio'].get('fiPct'); rec['kpi']['duration'] = o['portfolio'].get('duration')
        units = o.get('units', {})
        if units.get('total'): rec['kpi'].setdefault('units', units['total']['cur'])
        seg = o.get('segments', {})
        entry = {'id': qid, 'fy': r['fy'], 'q': r['q'], 'source': {'url': r.get('url'), 'date': r.get('date')},
                 'units': {k: v['cur'] for k, v in units.items()}, 'unitsPrevQ': {k: v['prevQ'] for k, v in units.items() if v.get('prevQ') is not None},
                 'unitsPrevY': {k: v['prevY'] for k, v in units.items() if v.get('prevY') is not None},
                 'premiums': {k: v['q'][0] * 1000 for k, v in seg.items()}, 'premiumsPrevY': {k: v['q'][1] * 1000 for k, v in seg.items()},
                 'premiumsYtd': {k: v['ytd'][0] * 1000 for k, v in seg.items() if v.get('ytd')}, 'premiumsYtdPrevY': {k: v['ytd'][1] * 1000 for k, v in seg.items() if v.get('ytd')},
                 'subsidiaries': {k: v['q'][0] * 1000 for k, v in o.get('subsidiaries', {}).items()}, 'subsidiariesPrevY': {k: v['q'][1] * 1000 for k, v in o.get('subsidiaries', {}).items()},
                 'subsidiariesYtd': {k: v['ytd'][0] * 1000 for k, v in o.get('subsidiaries', {}).items() if v.get('ytd')},
                 'solvency': o.get('solvency'), 'portfolio': o.get('portfolio'), 'reported': rep}
        OPS.append(entry)
    # comparative operating data for quarters without a report (prior-year columns of the following year's report)
    have = {e['id'] for e in OPS}
    for e in list(OPS):
        pq = prev_year_q(e['id'])
        if pq in have: continue
        fy, q = qid_parts(pq)
        OPS.append({'id': pq, 'fy': fy, 'q': q, 'source': dict(e['source'], comparative=True), 'units': dict(e['unitsPrevY']), 'unitsPrevQ': {}, 'unitsPrevY': {},
                    'premiums': dict(e['premiumsPrevY']), 'premiumsPrevY': {}, 'premiumsYtd': dict(e['premiumsYtdPrevY']), 'premiumsYtdPrevY': {},
                    'subsidiaries': dict(e['subsidiariesPrevY']), 'subsidiariesPrevY': {}, 'subsidiariesYtd': {}, 'solvency': None, 'portfolio': None, 'reported': {}, 'comparative': True})
        have.add(pq)
    OPS.sort(key=lambda e: (e['fy'], e['q']))
    for e in OPS:
        rec = Q.get(e['id'])
        if rec and e['units'].get('total') and 'units' not in rec['kpi']:
            rec['kpi']['units'] = e['units']['total']
    # strip private keys, sort, write
    def clean(d):
        return {k: v for k, v in d.items() if not k.startswith('_')}
    quarters = [dict(r, is_=None) for r in []]  # placeholder to keep linters quiet
    quarters = []
    for qid in sorted(Q, key=lambda x: qid_parts(x)):
        r = Q[qid]
        if not r['is'].get('written'): continue
        quarters.append({'id': r['id'], 'fy': r['fy'], 'q': r['q'], 'label': r['label'], 'is': clean(r['is']), 'bs': clean(r['bs']) or None, 'cf': clean(r['cf']) or None, 'kpi': r['kpi'], 'sources': r['sources'], 'shares': r['shares']})
    ytd = []
    for k in sorted(YTD, key=lambda x: (int(x[:4]), int(x.split('M')[1]))):
        r = YTD[k]
        if not r['is'].get('written') and not r['cf']: continue
        r['sources']['is'] = r['is'].get('_src'); r['sources']['cf'] = r['cf'].get('_src')
        ytd.append({'id': r['id'], 'fy': r['fy'], 'months': r['months'], 'is': clean(r['is']) or None, 'cf': clean(r['cf']) or None, 'kpi': r['kpi'], 'sources': {k2: v for k2, v in r['sources'].items() if v}})
    years = []
    for k in sorted(FY, key=lambda x: int(x[2:])):
        r = FY[k]
        if not r['is'].get('written'): continue
        q4 = Q.get('%dQ4' % r['fy'])
        r['sources']['is'] = r['is'].get('_src') or (q4 and q4['sources'].get('is')); r['sources']['cf'] = r['cf'].get('_src')
        years.append({'id': r['id'], 'fy': r['fy'], 'is': clean(r['is']), 'bs': (q4 and clean(q4['bs'])) or None, 'cf': clean(r['cf']) or None, 'kpi': r['kpi'], 'sources': {k2: v for k2, v in r['sources'].items() if v}, 'shares': (q4 and q4['shares']) or {}})
    layout = build_layout()
    gen = datetime.now(timezone.utc).isoformat(timespec='seconds')
    fin = {'generatedAt': gen, 'currency': 'MXN', 'units': 'thousands', 'unitsNote': 'Statements in thousands of pesos (CNSF format, as reported by Quálitas); ratios in %; insured units in thousands.',
           'layout': layout, 'quarters': quarters, 'ytd': ytd, 'years': years,
           'coverage': {'quarters': [quarters[0]['id'], quarters[-1]['id']] if quarters else [], 'years': [years[0]['id'], years[-1]['id']] if years else [], 'reportsParsed': len(reports), 'sificParsed': len(sifics)}}
    os.makedirs(OUT_DIR, exist_ok=True)
    hdr = '// AUTO-GENERATED by scripts/qualitas/build_data.py from tools/qualitas/raw — do not hand-edit.\n// Generated: %s\n' % gen
    with open(os.path.join(OUT_DIR, 'financials.js'), 'w', encoding='utf-8') as f:
        f.write(hdr + 'window.Q_FIN = ' + json.dumps(fin, ensure_ascii=False, separators=(',', ':')) + ';\n')
    ops_out = {'generatedAt': gen, 'units': 'insured units in thousands; premiums in thousands of pesos', 'quarters': OPS,
               'note': 'Written premiums by line of business and insured units by country as disclosed in each quarterly report; the four lines do not always add to the reported total (consolidation adjustments). Quarters marked comparative come from the prior-year column of the following year\'s report.'}
    with open(os.path.join(OUT_DIR, 'operations.js'), 'w', encoding='utf-8') as f:
        f.write(hdr + 'window.Q_OPS = ' + json.dumps(ops_out, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print('financials.js: %d quarters (%s → %s), %d YTD periods, %d fiscal years; operations.js: %d quarters' % (len(quarters), quarters[0]['id'] if quarters else '-', quarters[-1]['id'] if quarters else '-', len(ytd), len(years), len(OPS)))
    missing = [q['id'] for q in quarters if not q['cf']]
    print('quarters without a cash-flow statement:', ', '.join(missing[-12:]) if missing else 'none')
    for qid in missing[-12:]: logw('financials.js', '%s has no quarterly cash-flow statement (SIFIC filing missing or image-only)' % qid)
    with open(BUILD_LOG, 'w', encoding='utf-8') as f:
        json.dump({'generatedAt': gen, 'reportsParsed': len(reports), 'sificParsed': len(sifics), 'warnings': LOG}, f, ensure_ascii=False, indent=1)


def build_layout():
    L = lambda k, en, es, **kw: dict({'k': k, 'en': en, 'es': es}, **kw)
    is_rows = [
        L('written', 'Written premiums', 'Primas emitidas', level=0, bold=True),
        L('ceded', 'Ceded premiums', 'Primas cedidas', level=1),
        L('retained', 'Retained premiums', 'Primas de retención', level=0, bold=True),
        L('reserveInc', 'Net increase in unearned-premium reserve', 'Incremento neto de la reserva de riesgos en curso', level=1),
        L('earned', 'Earned premiums (net)', 'Primas de retención devengadas', level=0, bold=True),
        L('acqCost', 'Net acquisition cost', 'Costo neto de adquisición', level=1),
        L('commAgents', 'Commissions to agents', 'Comisiones a agentes', level=2),
        L('addlComp', 'Additional compensation to agents', 'Compensaciones adicionales a agentes', level=2),
        L('reinsCommTaken', 'Commissions on reinsurance assumed', 'Comisiones por reaseguro tomado', level=2),
        L('reinsCommCeded', 'Less: commissions on ceded reinsurance', '(−) Comisiones por reaseguro cedido', level=2),
        L('xlCover', 'Excess-of-loss reinsurance cover', 'Cobertura de exceso de pérdida', level=2),
        L('acqOther', 'Other acquisition costs (incl. payments to financial institutions)', 'Otros (incluye pagos a instituciones financieras, UDI)', level=2),
        L('lossCost', 'Net claims cost', 'Costo neto de siniestralidad', level=1),
        L('lossGross', 'Claims and other obligations', 'Siniestralidad y otras obligaciones pendientes de cumplir', level=2),
        L('lossRecovered', 'Less: recovered from non-proportional reinsurance', '(−) Siniestralidad recuperada del reaseguro no proporcional', level=2),
        L('techResult', 'Technical result', 'Utilidad (pérdida) técnica', level=0, bold=True),
        L('analogous', 'Result of similar and related operations', 'Resultado de operaciones análogas y conexas', level=1),
        L('grossProfit', 'Gross profit', 'Utilidad (pérdida) bruta', level=0, bold=True),
        L('opex', 'Net operating expenses', 'Gastos de operación netos', level=1),
        L('opexAdmin', 'Administrative and operating expenses (net of policy fees)', 'Gastos administrativos y operativos (netos de derechos de póliza)', level=2),
        L('opexStaff', 'Personnel', 'Remuneraciones y prestaciones al personal', level=2),
        L('opexDA', 'Depreciation and amortization', 'Depreciaciones y amortizaciones', level=2),
        L('opResult', 'Operating result', 'Utilidad (pérdida) de la operación', level=0, bold=True),
        L('rif', 'Comprehensive financing result (RIF)', 'Resultado integral de financiamiento (RIF)', level=1),
        L('rifInvestments', 'Investment income', 'De inversiones', level=2),
        L('rifSale', 'Gains (losses) on sale of investments', 'Por venta de inversiones', level=2),
        L('rifValuation', 'Valuation of investments', 'Por valuación de inversiones', level=2),
        L('rifSurcharge', 'Surcharge on premiums paid in instalments', 'Por recargo sobre primas', level=2),
        L('rifLoanInterest', 'Interest on loans', 'Intereses por créditos', level=2),
        L('rifWriteoffReins', 'Less: provisions on reinsurance recoverables', '(−) Castigos preventivos por importes recuperables de reaseguro', level=2),
        L('rifWriteoffCredit', 'Less: credit-risk provisions', '(−) Castigos preventivos por riesgos crediticios', level=2),
        L('rifOther', 'Other', 'Otros', level=2),
        L('rifFx', 'Foreign-exchange result', 'Resultado cambiario', level=2),
        L('associates', 'Share of results of permanent investments', 'Participación en el resultado de inversiones permanentes', level=1),
        L('ibt', 'Income before income taxes', 'Utilidad (pérdida) antes de impuestos a la utilidad', level=0, bold=True),
        L('tax', 'Income taxes', 'Provisión para el pago del impuesto a la utilidad', level=1),
        L('discontinued', 'Discontinued operations', 'Operaciones discontinuadas', level=1),
        L('netIncome', 'Net income', 'Utilidad (pérdida) del ejercicio', level=0, bold=True),
        L('nci', 'Non-controlling interest', 'Participación no controladora', level=1),
        L('netControlling', 'Net income attributable to controlling interest', 'Participación controladora', level=0, bold=True),
        # derived rows (kpi) — never additive
        L('acqRatio', 'Acquisition ratio (÷ retained premiums)', 'Índice de adquisición (÷ prima retenida)', level=1, kpi=True, pct=True),
        L('lossRatio', 'Loss ratio (÷ earned premiums)', 'Índice de siniestralidad (÷ prima devengada)', level=1, kpi=True, pct=True),
        L('opRatio', 'Operating ratio (÷ written premiums)', 'Índice de operación (÷ prima emitida)', level=1, kpi=True, pct=True),
        L('combined', 'Combined ratio', 'Índice combinado', level=0, bold=True, kpi=True, pct=True),
        L('combinedAdj', 'Adjusted combined ratio (all costs ÷ earned premiums)', 'Índice combinado ajustado (costos ÷ prima devengada)', level=1, kpi=True, pct=True),
        L('techMargin', 'Technical margin (÷ earned premiums)', 'Margen técnico (÷ prima devengada)', level=1, kpi=True, pct=True),
        L('opMargin', 'Operating margin (÷ earned premiums)', 'Margen operativo (÷ prima devengada)', level=1, kpi=True, pct=True),
        L('netMargin', 'Net margin (÷ written premiums)', 'Margen neto (÷ prima emitida)', level=1, kpi=True, pct=True),
        L('taxRate', 'Effective tax rate', 'Tasa efectiva de impuestos', level=1, kpi=True, pct=True),
        L('eps', 'Net income per share (Ps.)', 'Utilidad neta por acción (Ps.)', level=1, kpi=True, perShare=True),
    ]
    bs_rows = [
        L('inv', 'Investments', 'Inversiones', level=0, bold=True),
        L('securities', 'Securities', 'Valores', level=1),
        L('gov', 'Government securities', 'Gubernamentales', level=2), L('corpFixed', 'Corporate fixed-rate securities', 'Empresas privadas, tasa conocida', level=2),
        L('equities', 'Equities', 'Empresas privadas, renta variable', level=2), L('foreign', 'Foreign securities', 'Extranjeros', level=2),
        L('impairment', 'Less: impairment of securities', '(−) Deterioro de valores', level=2), L('restricted', 'Restricted securities', 'Valores restringidos', level=2),
        L('repo', 'Repo receivables', 'Deudor por reporto', level=1), L('loans', 'Loan portfolio, net', 'Cartera de crédito (neto)', level=1), L('realEstate', 'Real estate, net', 'Inmuebles (neto)', level=1),
        L('laborInv', 'Investments for labor obligations', 'Inversiones para obligaciones laborales', level=1),
        L('cash', 'Cash and banks', 'Disponibilidad (caja y bancos)', level=0, bold=True),
        L('receivables', 'Receivables', 'Deudores', level=0, bold=True), L('premRec', 'Premium receivables', 'Deudor por primas', level=2),
        L('govRec', 'Receivables from federal government entities', 'Adeudos a cargo de dependencias federales', level=2), L('agentsRec', 'Agents and adjusters', 'Agentes y ajustadores', level=2),
        L('otherRec', 'Other receivables', 'Otros deudores', level=2), L('recProvision', 'Less: allowance for write-offs', '(−) Estimación para castigos', level=2),
        L('reinsurers', 'Reinsurers, net', 'Reaseguradores y reafianzadores (neto)', level=1), L('permInv', 'Permanent investments', 'Inversiones permanentes', level=1),
        L('otherAssets', 'Other assets', 'Otros activos', level=1), L('equipment', 'Furniture and equipment, net', 'Mobiliario y equipo (neto)', level=2),
        L('otherAssetsMisc', 'Sundry', 'Diversos', level=2), L('intangAmort', 'Amortizable intangibles', 'Activos intangibles amortizables', level=2), L('intangLong', 'Long-lived intangibles', 'Activos intangibles de larga duración', level=2),
        L('totalAssets', 'Total assets', 'Suma del activo', level=0, bold=True),
        L('reserves', 'Technical reserves', 'Reservas técnicas', level=0, bold=True), L('upr', 'Unearned-premium (risks-in-force) reserve', 'Reserva de riesgos en curso', level=1),
        L('claimsReserves', 'Reserve for obligations pending settlement', 'Reserva para obligaciones pendientes de cumplir', level=1), L('claimsPending', 'Reported claims pending payment', 'Por pólizas vencidas y siniestros ocurridos pendientes de pago', level=2),
        L('ibnr', 'Incurred but not reported (IBNR) and adjustment expenses', 'Siniestros ocurridos no reportados y gastos de ajuste', level=2),
        L('fundsAdmin', 'Funds under administration', 'Fondos en administración', level=2), L('premDeposit', 'Premiums in deposit', 'Primas en depósito', level=2),
        L('laborReserves', 'Reserves for labor obligations', 'Reservas para obligaciones laborales', level=1),
        L('creditors', 'Creditors', 'Acreedores', level=0, bold=True), L('creditorsAgents', 'Agents and adjusters', 'Agentes y ajustadores', level=2), L('lossFunds', 'Loss funds under administration', 'Fondos en administración de pérdidas', level=2), L('creditorsMisc', 'Sundry creditors', 'Diversos', level=2),
        L('reinsPayable', 'Reinsurers', 'Reaseguradores y reafianzadores', level=1),
        L('otherLiab', 'Other liabilities', 'Otros pasivos', level=0, bold=True), L('ptuProv', 'Employee profit-sharing provision', 'Provisión para PTU', level=2), L('taxProv', 'Income-tax provision', 'Provisión para el pago de impuestos', level=2),
        L('otherObl', 'Other obligations', 'Otras obligaciones', level=2), L('deferredCredits', 'Deferred credits', 'Créditos diferidos', level=2),
        L('totalLiab', 'Total liabilities', 'Suma del pasivo', level=0, bold=True),
        L('paidCapital', 'Paid-in capital', 'Capital social pagado', level=1), L('treasury', 'Less: repurchased shares', '(−) Acciones propias recompradas', level=2),
        L('capReserves', 'Capital reserves', 'Reservas de capital', level=1), L('legalReserve', 'Legal reserve', 'Reserva legal', level=2), L('buybackReserve', 'Reserve for share repurchases', 'Reserva para adquisición de acciones propias', level=2), L('otherCapReserves', 'Other reserves', 'Otras reservas', level=2),
        L('valuationSurplus', 'Valuation surplus (available-for-sale securities)', 'Superávit por valuación', level=1), L('retained', 'Retained earnings', 'Resultados de ejercicios anteriores', level=1),
        L('netIncomeEq', 'Net income for the period', 'Resultado del ejercicio', level=1), L('controllingEquity', 'Controlling interest', 'Participación controladora', level=0, bold=True),
        L('nci', 'Non-controlling interest', 'Participación no controladora', level=1), L('totalEquity', "Total stockholders' equity", 'Suma del capital contable', level=0, bold=True),
        L('totalLiabEquity', "Total liabilities and stockholders' equity", 'Suma del pasivo y capital contable', level=0, bold=True),
    ]
    cf_rows = [
        L('netIncome', 'Net income', 'Resultado neto', level=0, bold=True),
        L('nonCashAdj', 'Adjustments for non-cash items', 'Ajustes por partidas que no implican flujo de efectivo', level=1),
        L('valuation', 'Valuation (gains) losses on investments', 'Utilidad o pérdida por valorización de inversiones', level=2), L('badDebt', 'Allowance for doubtful accounts', 'Estimación para castigo o difícil cobro', level=2),
        L('da', 'Depreciation and amortization', 'Depreciaciones y amortizaciones', level=2), L('reservesAdj', 'Increase in technical reserves', 'Ajuste o incremento a las reservas técnicas', level=2),
        L('provisions', 'Provisions', 'Provisiones', level=2), L('taxes', 'Current and deferred income taxes', 'Impuestos a la utilidad causados y diferidos', level=2),
        L('chgSecurities', 'Change in securities', 'Cambio en inversiones en valores', level=1), L('chgRepo', 'Change in repo receivables', 'Cambio en deudores por reporto', level=1),
        L('chgPremRec', 'Change in premium receivables', 'Cambio en primas por cobrar', level=1), L('chgReceivables', 'Change in other receivables', 'Cambio en deudores', level=1),
        L('chgReinsurers', 'Change in reinsurers', 'Cambio en reaseguradoras y reafianzadoras', level=1), L('chgOtherOpAssets', 'Change in other operating assets', 'Cambio en otros activos operativos', level=1),
        L('chgClaimsObligations', 'Change in contractual obligations and claims expenses', 'Cambio en obligaciones contractuales y gastos asociados a la siniestralidad', level=1),
        L('chgOtherOpLiab', 'Change in other operating liabilities', 'Cambio en otros pasivos operativos', level=1),
        L('opChanges', 'Net change in operating items (CNSF subtotal, as filed)', 'Flujos netos de actividades de operación (subtotal CNSF, según SIFIC)', level=1),
        L('cfo', 'Net cash from operating activities (net income + adjustments + operating items)', 'Flujo de efectivo de operación (resultado neto + ajustes + partidas operativas)', level=0, bold=True),
        L('saleProceedsPPE', 'Proceeds from disposal of property and equipment', 'Cobros por disposición de inmuebles, mobiliario y equipo', level=1), L('capexPPE', 'Purchases of property and equipment', 'Pagos por adquisición de inmuebles, mobiliario y equipo', level=1),
        L('acqSubs', 'Acquisition of subsidiaries and associates', 'Pagos por adquisición de subsidiarias y asociadas', level=1), L('acqPermInv', 'Acquisition of other permanent investments', 'Pagos por adquisición de otras inversiones permanentes', level=1),
        L('intangibles', 'Purchases of intangible assets', 'Pagos por adquisición de activos intangibles', level=1),
        L('cfi', 'Net cash from investing activities', 'Flujos netos de efectivo de actividades de inversión', level=0, bold=True),
        L('dividendsPaid', 'Dividends paid', 'Pagos de dividendos en efectivo', level=1), L('buybacks', 'Share repurchases, net', 'Pagos asociados a la recompra de acciones propias', level=1),
        L('cff', 'Net cash from financing activities', 'Flujos netos de efectivo de actividades de financiamiento', level=0, bold=True),
        L('netChangeCash', 'Net increase (decrease) in cash', 'Incremento o disminución neta de efectivo', level=0, bold=True),
        L('fxEffectCash', 'Effect of exchange-rate changes on cash', 'Efectos por cambios en el valor del efectivo', level=1),
        L('cashBegin', 'Cash at beginning of period', 'Efectivo al inicio del periodo', level=1), L('cashEnd', 'Cash at end of period', 'Efectivo al final del periodo', level=0, bold=True),
    ]
    kpi_rows = [
        L('units', 'Insured units (thousands)', 'Unidades aseguradas (miles)'), L('rsi', 'Return on investments, RSI (%, as reported)', 'Rendimiento sobre inversiones, RSI (%, reportado)'),
        L('roe12', 'ROE, trailing 12 months (%, as reported)', 'ROE 12 meses (%, reportado)'), L('roePeriod', 'ROE for the period, annualized (%, as reported)', 'ROE del periodo, anualizado (%, reportado)'),
        L('rcs', 'Regulatory capital requirement (RCS)', 'Requerimiento de capital de solvencia (RCS)'), L('solvMargin', 'Solvency margin', 'Margen de solvencia'), L('solvIndex', 'Solvency index (%)', 'Índice de solvencia (%)'),
        L('fiPct', 'Fixed income, % of portfolio', 'Renta fija, % del portafolio'), L('duration', 'Portfolio duration (years)', 'Duración del portafolio (años)'), L('float', 'Invested assets / float (Ps. M, as reported)', 'Activos invertidos / float (Ps. M, reportado)'),
    ]
    return {'is': is_rows, 'bs': bs_rows, 'cf': cf_rows, 'kpi': kpi_rows}


if __name__ == '__main__':
    main()
