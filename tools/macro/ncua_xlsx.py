"""Reads the NCUA's quarterly "Financial Trends in Federally Insured Credit Unions" chart
pack workbook (the xlsx inside chart-pack-YYYY-qN.zip; forty quarters per edition) and
writes the series the dashboard uses as JSON.

    python ncua_xlsx.py <workbook.xlsx> <out.json>

Sheets are found by the chart title in cell A1, not by number:
  "Number of Insured Credit Unions Reporting"      -> count (FCU + FISCU = "Both")
  "Aggregated Net Worth Ratio"                      -> net worth / total assets, all FICUs
  "Delinquency & Net Charge-Offs"                   -> delinquency rate (loans 60+ days)
  "Borrowings to Total Shares and Net Worth"        -> borrowings ratio
  "Return on Average Assets vs. Provision"          -> ROAA (annualized)
  "Summary of Trends by CU Type"                    -> total assets, loans and shares in dollars ("ALL"),
                                                       and the net worth ratio the pack prints, cross-checked
Ratios are fractions in the file and written in percent. Quarters are "2016Q3" in the file
and written as "2016-Q3". The workbook's own summary net worth ratio must agree with the
chart series (rounded to one decimal) or the script fails.
"""
import json, re, sys
import openpyxl

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

def sheet(title_re):
    for ws in wb.worksheets:
        a1 = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if a1 and a1[0] and re.search(title_re, str(a1[0]), re.I): return ws
    raise SystemExit("chart not found: " + title_re)

def quarter(v):
    m = re.match(r'^\s*(\d{4})Q([1-4])\s*$', str(v))
    return "%s-Q%s" % (m.group(1), m.group(2)) if m else None

def series(ws, col_re):
    """{quarter: value} of the column whose header (row 3) matches col_re."""
    rows = list(ws.iter_rows(values_only=True))
    hdr = [re.sub(r'\s+', ' ', str(v)).strip() if v is not None else "" for v in rows[2]]
    if hdr[0].lower() != "date": raise SystemExit("no Date header in " + ws.title)
    idx = [i for i, h in enumerate(hdr) if re.search(col_re, h, re.I)]
    if len(idx) != 1: raise SystemExit("column /%s/ matches %d headers in %s: %s" % (col_re, len(idx), ws.title, hdr))
    outd = {}
    for r in rows[3:]:
        q = quarter(r[0]) if r and r[0] is not None else None
        if not q: continue
        v = r[idx[0]] if idx[0] < len(r) else None
        outd[q] = float(v) if isinstance(v, (int, float)) else None
    if len(outd) < 36: raise SystemExit("too few quarters in " + ws.title)
    return outd

count = series(sheet(r'^Chart \d+: Number of Insured Credit Unions Reporting'), r'^Both$')
nw = series(sheet(r'^Chart \d+: Aggregated Net Worth Ratio'), r'^Net Worth Ratio$')
delq = series(sheet(r'^Chart \d+: Delinquency & Net Charge-Offs'), r'^Delinquency Rate$')
borrow = series(sheet(r'^Chart \d+: Borrowings to Total Shares and Net Worth'), r'^Borrowings to Total Share & Net Worth$')
roaa = series(sheet(r'^Chart \d+: Return on Average Assets vs\. Provision'), r'^Return on Average Assets \(Annualized\)$')

quarters = sorted(nw)
for a, b in zip(quarters, quarters[1:]):
    ya, qa = int(a[:4]), int(a[6]); yb, qb = int(b[:4]), int(b[6])
    if yb * 4 + qb != ya * 4 + qa + 1: raise SystemExit("quarters are not consecutive: %s -> %s" % (a, b))
pc = lambda v, d=2: None if v is None else round(v * 100, d)
quarterly = []
for q in quarters:
    if nw[q] is None or count.get(q) is None: raise SystemExit("missing net worth ratio or count in " + q)
    if not 0.05 <= nw[q] <= 0.20: raise SystemExit("implausible net worth ratio in %s: %r" % (q, nw[q]))
    if not 2500 <= count[q] <= 8000: raise SystemExit("implausible credit union count in %s: %r" % (q, count[q]))
    quarterly.append({"d": q, "n": int(count[q]), "nw": pc(nw[q]), "delq": pc(delq.get(q)), "borrow": pc(borrow.get(q)), "roaa": pc(roaa.get(q))})

# The summary sheet: dollar totals and the net worth ratio as printed, for the cross-check.
summ = sheet(r'^Chart \d+: Summary of Trends by CU Type')
rows = list(summ.iter_rows(values_only=True))
hdr = [str(v).strip() if v is not None else "" for v in rows[2]]
if "ALL" not in hdr: raise SystemExit("no ALL column in the summary sheet")
col = hdr.index("ALL")
def cell(label_re):
    for r in rows[3:]:
        if r and r[0] is not None and re.match(label_re, str(r[0]).strip(), re.I): return r[col]
    raise SystemExit("summary row /%s/ not found" % label_re)
def dollars(v): return round(float(re.sub(r'[^\d.]', '', str(v))) / 1e9, 1)   # "$2,498,852,640,655" -> 2498.9 ($ billions)
assets, loans, shares = dollars(cell(r'^Total Assets$')), dollars(cell(r'^Total Loans$')), dollars(cell(r'^Total Shares$'))
nCheck = int(cell(r'^Number of FICU Reporting$'))
nwPrinted = float(re.sub(r'[^\d.]', '', str(cell(r'^Net Worth Ratio$'))))
last = quarterly[-1]
if nCheck != last["n"]: raise SystemExit("summary count %d differs from chart count %d" % (nCheck, last["n"]))
if abs(nwPrinted - round(last["nw"], 1)) > 0.051: raise SystemExit("summary net worth ratio %.1f%% differs from the chart series %.2f%%" % (nwPrinted, last["nw"]))
if not 500 <= assets <= 10000: raise SystemExit("implausible total assets %r" % assets)
json.dump({"quarterly": quarterly, "asOf": last["d"], "from": quarterly[0]["d"], "assets": assets, "loans": loans, "shares": shares}, open(out, "w"), separators=(",", ":"))
print("NCUA chart pack: %d quarters %s .. %s; %d credit unions, assets $%.1f bn, net worth ratio %.2f%%, delinquency %.2f%%, borrowings %.2f%% of shares and net worth, ROAA %.2f%%" % (
    len(quarterly), quarterly[0]["d"], last["d"], last["n"], assets, last["nw"], last["delq"], last["borrow"], last["roaa"]))
