"""Reads the New York Fed's Quarterly Report on Household Debt and Credit data workbook
(HHD_C_Report_YYYYQn.xlsx) and writes the series the dashboard uses as JSON.

    python hhdc_xlsx.py <workbook.xlsx> <out.json>

Sheets are found by the title in their first rows, not by name or position, because the
report's page numbers move between editions:
  "Total Debt Balance and Its Composition"            -> balances by loan type, $ trillions
  "Percent of Balance 90+ Days Delinquent by Loan Type" -> share of each type's balance 90+ days late
  "Total Balance by Delinquency Status"               -> checked (shares sum to 100), not written
Quarters are labelled "03:Q1" in the file and written as "2003-Q1". Every balance row must
add up to its total (within $0.02 trillion) or the script fails, so a layout change can
never ship wrong figures.
"""
import json, re, sys
import openpyxl

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

def quarter(label):
    m = re.match(r'^\s*(\d{2}):Q([1-4])\s*$', str(label))
    if not m: return None
    yy = int(m.group(1)); return "%d-Q%s" % (1900 + yy if yy > 50 else 2000 + yy, m.group(2))

def find_sheet(title_re):
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
            if any(v is not None and re.search(title_re, str(v), re.I) for v in row): return ws
    raise SystemExit("sheet not found: " + title_re)

def table(ws, headers):
    """Rows of the sheet whose first cell is a quarter label; columns named by `headers`
    (the header row of the sheet is read for the order, matched case-insensitively)."""
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    for r in rows[:8]:
        cells = [str(v).strip().lower() if v is not None else "" for v in r]
        if sum(1 for h in headers if any(h.lower() in c for c in cells)) >= len(headers) - 1: hdr = cells; break
    if hdr is None: raise SystemExit("header row not found in " + ws.title)
    cols = []
    for h in headers:
        idx = next((i for i, c in enumerate(hdr) if h.lower() in c), None)
        if idx is None:
            # The first header cell of some sheets is a stray number; the first data column is then column 1.
            if h.lower() in ("mortgage",): idx = 1
            else: raise SystemExit("column %s not found in %s" % (h, ws.title))
        cols.append(idx)
    out_rows = []
    for r in rows:
        q = quarter(r[0]) if r and r[0] is not None else None
        if not q: continue
        vals = []
        for i in cols:
            v = r[i] if i < len(r) else None
            vals.append(round(float(v), 4) if isinstance(v, (int, float)) else None)
        out_rows.append([q] + vals)
    if len(out_rows) < 40: raise SystemExit("too few rows in " + ws.title)
    return out_rows

bal_cols = ["Mortgage", "HE Revolving", "Auto Loan", "Credit Card", "Student Loan", "Other", "Total"]
bal = table(find_sheet(r'Total Debt Balance and Its Composition'), bal_cols)
keys = ["mortgage", "heloc", "auto", "card", "student", "other", "total"]
balances = []
for r in bal:
    row = dict(zip(["d"] + keys, r))
    if any(row[k] is None for k in keys): raise SystemExit("missing balance in " + r[0])
    parts = sum(row[k] for k in keys[:-1])
    if abs(parts - row["total"]) > 0.02: raise SystemExit("balances do not add up in %s: %.4f vs %.4f" % (r[0], parts, row["total"]))
    balances.append(row)

dq_cols = ["Mortgage", "HELOC", "Auto", "CC", "Student Loan", "Other", "All"]
dq = table(find_sheet(r'Percent of Balance 90\+ Days Delinquent by Loan Type'), dq_cols)
delinq90 = [dict(zip(["d", "mortgage", "heloc", "auto", "card", "student", "other", "all"], [r[0]] + [None if v is None else round(v, 2) for v in r[1:]])) for r in dq]

st_cols = ["Current", "30 days late", "60 days late", "90 days late", "120+ days late", "Severely Derogatory"]
st = table(find_sheet(r'Total Balance by Delinquency Status'), st_cols)
status = [dict(zip(["d", "current", "d30", "d60", "d90", "d120", "derog"], [r[0]] + [None if v is None else round(v, 2) for v in r[1:]])) for r in st]
for row in status:
    s = sum(v for k, v in row.items() if k != "d" and v is not None)
    if abs(s - 100) > 0.5: raise SystemExit("delinquency status shares do not sum to 100 in " + row["d"])

if balances[-1]["d"] != delinq90[-1]["d"]: raise SystemExit("balances end %s but delinquency ends %s" % (balances[-1]["d"], delinq90[-1]["d"]))
json.dump({"asOf": balances[-1]["d"], "balances": balances, "delinq90": delinq90}, open(out, "w"), separators=(",", ":"))
print("hhdc: %d quarters %s .. %s; total %.3f tn, mortgage %.3f, card %.3f; 90+ delinquent %.2f%%" % (len(balances), balances[0]["d"], balances[-1]["d"], balances[-1]["total"], balances[-1]["mortgage"], balances[-1]["card"], delinq90[-1]["all"]))
