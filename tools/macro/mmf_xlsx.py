"""Reads the SEC's Money Market Fund Statistics "supporting data" workbook (Form N-MFP
aggregates, monthly since December 2010) and writes the series the dashboard uses as JSON.

    python mmf_xlsx.py <workbook.xlsx> <out.json>

Tabs are found by their title in the Notes sheet, not by position:
  "MMF Net Assets ($Billions)"                                   -> net assets by category and fund type
  "Number of MMFs Reporting"                                     -> fund counts (total)
  "Aggregate MMF Daily Liquid Assets (Percent of Total Assets)"  -> daily liquid assets by fund type, since Oct-2016
  "Aggregate MMF Weekly Liquid Assets (Percent of Total Assets)" -> weekly liquid assets by fund type, since Oct-2016
Every month's category subtotals must equal the sum of their fund types (where the split
exists) and the total must equal the sum of the subtotals, within $0.5 billion, or the
script fails: a layout change can never ship wrong figures. Before October 2016 the prime
and tax-exempt retail/institutional split did not exist (the fund-type cells hold 0 and the
subtotal carries the figure); those cells are written as null.
"""
import json, re, sys, datetime
import openpyxl

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

def tab_for(title_re):
    """The 'Tab.x.y' sheet whose description in the Notes sheet matches title_re."""
    notes = wb["Notes"]
    for row in notes.iter_rows(values_only=True):
        if row and row[0] and re.match(r'^Tab\.\d+\.\d+$', str(row[0]).strip()) and len(row) > 1 and row[1] and re.search(title_re, str(row[1]).split("\n")[0], re.I):
            name = str(row[0]).strip()
            if name not in wb.sheetnames: raise SystemExit("sheet %s not in workbook" % name)
            return wb[name]
    raise SystemExit("no tab titled /%s/ in Notes" % title_re)

def month(v):
    if isinstance(v, datetime.datetime): return "%04d-%02d" % (v.year, v.month)
    m = re.match(r'^(\d{4})-(\d{2})', str(v))
    return m.group(1) + "-" + m.group(2) if m else None

def table(ws):
    """{ (category, fund type): {month: value} } plus the ordered month list."""
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    if not hdr or str(hdr[0]).strip().lower() != "category" or str(hdr[1]).strip().lower() != "fund type": raise SystemExit("unexpected header in " + ws.title)
    months = [month(v) for v in hdr[2:]]
    cols = [(i + 2, m) for i, m in enumerate(months) if m]
    if len(cols) < 100: raise SystemExit("too few months in " + ws.title)
    for a, b in zip(cols, cols[1:]):
        ya, ma = map(int, a[1].split("-")); yb, mb = map(int, b[1].split("-"))
        if yb * 12 + mb != ya * 12 + ma + 1: raise SystemExit("months are not consecutive in %s: %s -> %s" % (ws.title, a[1], b[1]))
    data = {}
    for r in rows[1:]:
        if not r or r[0] is None or str(r[0]).startswith("Return"): continue
        key = (str(r[0]).strip(), str(r[1]).strip() if r[1] is not None else "")
        data[key] = {m: (float(r[i]) if isinstance(r[i], (int, float)) else None) for i, m in cols}
    return data, [m for _, m in cols]

net, months = table(tab_for(r'^MMF Net Assets \(\$Billions\)'))
cnt, cmonths = table(tab_for(r'^Number of MMFs Reporting'))
dla, lmonths = table(tab_for(r'^Aggregate MMF Daily Liquid Assets'))
wla, wmonths = table(tab_for(r'^Aggregate MMF Weekly Liquid Assets'))
if cmonths != months: raise SystemExit("fund counts cover different months than net assets")
if lmonths[-1] != months[-1] or wmonths[-1] != months[-1]: raise SystemExit("liquidity tabs end at a different month")

def need(d, cat, ft):
    if (cat, ft) not in d: raise SystemExit("row %s / %s missing" % (cat, ft))
    return d[(cat, ft)]
gov, tsy, govSub = need(net, "Government", "Government"), need(net, "Government", "Treasury"), need(net, "Government", "Subtotal")
primeI, primeR, primeSub = need(net, "Prime", "Institutional"), need(net, "Prime", "Retail"), need(net, "Prime", "Subtotal")
texI, texR, texSub = need(net, "Tax Exempt", "Institutional"), need(net, "Tax Exempt", "Retail"), need(net, "Tax Exempt", "Subtotal")
total = need(net, "Total", "Total")
nTotal = need(cnt, "Total", "Total")

def split(a, b, sub, m):
    """Institutional/retail cells, null where the split did not exist (both 0 under a nonzero subtotal)."""
    x, y, s = a[m], b[m], sub[m]
    if x is None or y is None or s is None: return None, None
    if x == 0 and y == 0 and s > 0: return None, None
    if abs(x + y - s) > 0.5: raise SystemExit("fund types do not add up to the subtotal in %s: %.2f + %.2f vs %.2f" % (m, x, y, s))
    return round(x, 1), round(y, 1)

def pct(d, cat, ft, m):
    v = d.get((cat, ft), {}).get(m)
    if v is None: return None
    if v < 0 or v > 1.05: raise SystemExit("liquidity share out of range in %s / %s %s: %r" % (cat, ft, m, v))
    return round(v * 100, 1)

monthly = []
for m in months:
    g, t, gs, ps, ts, tt = gov[m], tsy[m], govSub[m], primeSub[m], texSub[m], total[m]
    if any(v is None for v in (g, t, gs, ps, ts, tt)): raise SystemExit("missing net assets in " + m)
    if abs(g + t - gs) > 0.5: raise SystemExit("government subtotal does not add up in " + m)
    if abs(gs + ps + ts - tt) > 0.5: raise SystemExit("total does not add up in %s: %.2f vs %.2f" % (m, gs + ps + ts, tt))
    pI, pR = split(primeI, primeR, primeSub, m)
    xI, xR = split(texI, texR, texSub, m)
    monthly.append({
        "d": m, "gov": round(g, 1), "tsy": round(t, 1), "primeI": pI, "primeR": pR, "prime": round(ps, 1), "texI": xI, "texR": xR, "tex": round(ts, 1), "total": round(tt, 1),
        "n": int(nTotal[m]) if nTotal.get(m) is not None else None,
        "dlaGov": pct(dla, "Government", "Government", m), "dlaTsy": pct(dla, "Government", "Treasury", m), "dlaPrimeI": pct(dla, "Prime", "Institutional", m), "dlaPrimeR": pct(dla, "Prime", "Retail", m),
        "wlaGov": pct(wla, "Government", "Government", m), "wlaTsy": pct(wla, "Government", "Treasury", m), "wlaPrimeI": pct(wla, "Prime", "Institutional", m), "wlaPrimeR": pct(wla, "Prime", "Retail", m),
        "wlaTexI": pct(wla, "Tax Exempt", "Institutional", m), "wlaTexR": pct(wla, "Tax Exempt", "Retail", m)
    })
if len(monthly) < 150: raise SystemExit("too few months: %d" % len(monthly))
last = monthly[-1]
if last["total"] < 2000 or last["total"] > 30000: raise SystemExit("implausible total net assets %r" % last["total"])
if last["wlaGov"] is None or last["wlaPrimeI"] is None: raise SystemExit("weekly liquid assets missing for " + last["d"])
json.dump({"monthly": monthly, "asOf": last["d"], "from": monthly[0]["d"], "liquidityFrom": next(r["d"] for r in monthly if r["wlaGov"] is not None)}, open(out, "w"), separators=(",", ":"))
print("MMF statistics: %d months %s .. %s; net assets $%.1f bn (government %.1f, treasury %.1f, prime %.1f, tax-exempt %.1f), %s funds; weekly liquid assets: government %.1f%%, prime institutional %.1f%%" % (
    len(monthly), monthly[0]["d"], last["d"], last["total"], last["gov"], last["tsy"], last["prime"], last["tex"], last["n"], last["wlaGov"], last["wlaPrimeI"]))
